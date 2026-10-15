#!/usr/bin/env python3
"""Make an image-embedded, manually annotatable webcam blink review workbook.

Run on the user's Mac with Empra9/Empra10 mounted. Only reads existing data;
never changes source gaze, timestamps, pipeline outputs or videos.
"""
import argparse
import base64
import io
import json
import subprocess
import sys
from pathlib import Path

import cv2
import mediapipe as mp
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

VPS = ['beo7','bjs4','egf5','fbn6','fgt6','jkl7','kdn8','kro3','mhe9','oem4','ogt7']
# Must match the original diagnostic image landmark ordering.
LEFT = [33,160,158,133,153,144]
RIGHT = [362,385,387,263,373,380]


def read_trials(path):
    obj = json.loads(path.read_text(encoding='utf-8'))['phases']
    items = obj.get('experiment_block1',{}).get('trials',[]) + obj.get('experiment_block2',{}).get('trials',[])
    out = {}
    for row in items:
        if row.get('is_practice', False): continue
        out[int(row['trial_number'])] = (
            float(row['fixation']['start_video_s'])*1000.,
            float(row['stimulus']['end_video_s'])*1000.)
    return out


def find_run(vpdir):
    required = ('phases_detected.json','debug_5_pupil_data_calibrated.csv',
                'debug_5_ptgaze_calibrated.csv','debug_3_eyetracker_data.csv')
    runs = [r for r in (vpdir/'Analyse').glob('Run_60hz_FullCalib_*')
            if r.is_dir() and all((r/f).is_file() for f in required)]
    if not runs: raise FileNotFoundError(f'No complete FullCalib run for {vpdir.name}')
    return max(runs, key=lambda r:r.stat().st_mtime)


def read_eye_events(path):
    d=pd.read_csv(path)
    return d.assign(
        trial_num=pd.to_numeric(d['trial'],errors='coerce'),
        mid_ms=(pd.to_numeric(d['start_from_fixation_ms'],errors='coerce')+
                pd.to_numeric(d['end_from_fixation_ms'],errors='coerce'))/2.)


def read_web_regions(path, groups):
    d=pd.read_csv(path)
    d=d[d['reason_group'].astype(str).isin(groups)].copy()
    for c in ('start_video_time_ms','end_video_time_ms','start_frame','end_frame'):
        d[c]=pd.to_numeric(d[c],errors='coerce')
    return d.dropna(subset=['start_video_time_ms','end_video_time_ms'])


def min_ear_frame(frame_df, lo_ms, hi_ms):
    t=pd.to_numeric(frame_df['video_time_ms'],errors='coerce')
    sample=frame_df[t.between(lo_ms,hi_ms,inclusive='both')].copy()
    if sample.empty: return None
    raw=sample[['left_ear_raw','right_ear_raw']].apply(pd.to_numeric,errors='coerce').mean(axis=1)
    return sample.loc[raw.idxmin()] if raw.notna().any() else sample.iloc[len(sample)//2]


def extract_frame_at_pts(video, time_s):
    cmd=['ffmpeg','-v','error','-ss',f'{max(float(time_s),0):.6f}',
         '-i',str(video),'-frames:v','1','-f','image2pipe','-vcodec','png','-']
    r=subprocess.run(cmd,capture_output=True)
    if r.returncode!=0 or not r.stdout: return None
    return cv2.imdecode(np.frombuffer(r.stdout,dtype=np.uint8),cv2.IMREAD_COLOR)


def eye_crop(mesh, frame):
    rgb=cv2.cvtColor(frame,cv2.COLOR_BGR2RGB)
    res=mesh.process(rgb)
    if not res.multi_face_landmarks:
        return cv2.resize(rgb,(360,116),interpolation=cv2.INTER_AREA)
    h,w=rgb.shape[:2]; lm=res.multi_face_landmarks[0]
    lp=np.array([[lm.landmark[i].x*w,lm.landmark[i].y*h] for i in LEFT])
    rp=np.array([[lm.landmark[i].x*w,lm.landmark[i].y*h] for i in RIGHT])
    p=np.vstack([lp,rp])
    x0=max(0,int(p[:,0].min()-55));x1=min(w,int(p[:,0].max()+55))
    y0=max(0,int(p[:,1].min()-45));y1=min(h,int(p[:,1].max()+45))
    crop=rgb[y0:y1,x0:x1].copy()
    if crop.size==0: return cv2.resize(rgb,(360,116),interpolation=cv2.INTER_AREA)
    for pts in (lp,rp):
        for k,(x,y) in enumerate(pts):
            cv2.circle(crop,(int(x-x0),int(y-y0)),4,(255,0,0),-1)
            cv2.putText(crop,str(k),(int(x-x0)+3,int(y-y0)-3),cv2.FONT_HERSHEY_SIMPLEX,.33,(255,255,0),1)
    return crop


def image_data_url(image_rgb):
    h,w=image_rgb.shape[:2]
    width=350
    height=max(60,round(h*width/max(w,1)))
    scaled=cv2.resize(image_rgb,(width,height),interpolation=cv2.INTER_AREA)
    ok,out=cv2.imencode('.png',cv2.cvtColor(scaled,cv2.COLOR_RGB2BGR))
    if not ok: return None
    return 'data:image/png;base64,'+base64.b64encode(out.tobytes()).decode('ascii')


def gather(args):
    rows=[]; pictures=[]; problems=[]
    for vp in args.vp or VPS:
        try:
            vpdir=args.results_root/vp; run=find_run(vpdir)
            test=vpdir/'test60'
            frame=pd.read_csv(test/'blink_filter_frame_log.csv')
            regions=read_web_regions(test/'blink_filter_regions.csv',set(args.reason_group))
            eyelink=read_eye_events(test/'eyelink_blink_benchmark.csv')
            trials=read_trials(run/'phases_detected.json')
            video=args.video_root/f'{vp}.mp4'
            if not video.is_file(): raise FileNotFoundError(video)
            print(f'[SOURCE] {vp}: {run.name}, {len(regions)} candidate regions',flush=True)
            with mp.solutions.face_mesh.FaceMesh(max_num_faces=1,refine_landmarks=True,
                  min_detection_confidence=.5,min_tracking_confidence=.5) as mesh:
                for tr,(start,end) in sorted(trials.items()):
                    if args.trial and tr not in args.trial: continue
                    trows=regions[(regions['end_video_time_ms']>start)&
                                  (regions['start_video_time_ms']<end)]
                    midpoints=eyelink[eyelink['trial_num']==tr]['mid_ms'].dropna().to_numpy(float)
                    for _,r in trows.iterrows():
                        # A region can cross trial edges. Use its clipped trial interval consistently.
                        lo=max(start,float(r['start_video_time_ms']))
                        hi=min(end,float(r['end_video_time_ms']))
                        if hi<=lo: continue
                        point=min_ear_frame(frame,lo,hi)
                        if point is None:
                            problems.append(f'{vp} trial {tr}: no frame inside [{lo},{hi}]')
                            continue
                        fnum=int(point['frame_number']); timems=float(point['video_time_ms'])
                        y=int(np.any((midpoints>=lo-start)&(midpoints<=hi-start)))
                        im=extract_frame_at_pts(video,timems/1000.)
                        pic=None if im is None else image_data_url(eye_crop(mesh,im))
                        if pic is None: problems.append(f'{vp} trial {tr}: frame {fnum} image unavailable')
                        rows.append([len(rows)+1,vp,tr,fnum,round(timems,2),y,None,round((lo-start)/1000,3),round((hi-start)/1000,3)])
                        pictures.append(pic)
        except Exception as exc:
            problems.append(f'{vp}: {type(exc).__name__}: {exc}')
            print(f'[ERROR] {problems[-1]}',file=sys.stderr)
    return rows,pictures,problems


def make_book(rows, pictures, problems, args):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Blink Review'
    summary = wb.create_sheet('Summary')
    ws.merge_cells('A1:J1')
    ws['A1'] = 'Webcam Blink — Visual Inspection'
    ws.merge_cells('A2:J2')
    ws['A2'] = 'Visual inspection: 1 = visually confirmed blink; 0 = not a blink. Blank = not yet reviewed.'
    ws.merge_cells('A3:J3')
    ws['A3'] = 'EyeLink: 1 iff an EBLINK midpoint falls inside the webcam region (no tolerance, no 40-ms filter).'
    headings = ['ID','VP','Trial','Min-EAR frame','Video time (ms)','EyeLink (0/1)',
                'Visual (0/1)','Web start (s)','Web end (s)','Eye ROI + landmarks']
    for col, value in enumerate(headings, 1):
        ws.cell(5, col, value)

    for i, row in enumerate(rows, 6):
        for col, value in enumerate(row, 1):
            if value is not None:
                ws.cell(i, col, value)
        ws.row_dimensions[i].height = 83
        ws.cell(i, 7).fill = PatternFill('solid', fgColor='FFF4CC')
        if pictures[i - 6]:
            raw = base64.b64decode(pictures[i - 6].split(',', 1)[1])
            pic = ExcelImage(io.BytesIO(raw))
            pic.width, pic.height = 285, 88
            ws.add_image(pic, f'J{i}')

    if rows:
        validation = DataValidation(type='list', formula1='"0,1"', allow_blank=True)
        validation.error = 'Enter 0 or 1'; validation.errorTitle = 'Invalid label'
        validation.showErrorMessage = True
        ws.add_data_validation(validation)
        validation.add(f'G6:G{5+len(rows)}')
        ws.conditional_formatting.add(f'G6:G{5+len(rows)}',
            CellIsRule(operator='equal', formula=['0'], fill=PatternFill('solid', fgColor='FEE2E2')))
        ws.conditional_formatting.add(f'G6:G{5+len(rows)}',
            CellIsRule(operator='equal', formula=['1'], fill=PatternFill('solid', fgColor='DCFCE7')))
        ws.auto_filter.ref = f'A5:J{5+len(rows)}'

    for col, width in {'A':8,'B':12,'C':10,'D':19,'E':20,'F':19,
                       'G':19,'H':18,'I':18,'J':49}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = 'D6'
    ws.row_dimensions[1].height = 31
    ws.row_dimensions[2].height = 23
    ws.row_dimensions[3].height = 23
    ws.row_dimensions[5].height = 30
    for rowno, bg in [(1, '102A43'), (2, 'E8F1F8'), (3, 'E8F1F8'), (5, '254D70')]:
        for cell in ws[rowno][:10]:
            cell.fill = PatternFill('solid', fgColor=bg)
            cell.font = Font(bold=rowno in (1, 5), color='FFFFFF' if rowno in (1, 5) else '263445',
                             size=16 if rowno == 1 else 10)
            cell.alignment = Alignment(vertical='center')
    for rowno in range(6, 6+len(rows)):
        for c in range(1, 10):
            ws.cell(rowno, c).alignment = Alignment(vertical='center', horizontal='center')

    summary.merge_cells('A1:E1')
    summary['A1'] = 'Visual Blink Confirmation — Summary'
    summary.merge_cells('A2:E2')
    summary['A2'] = 'Accuracy shown only after ALL events for that VP are reviewed.'
    for col, name in enumerate(['VP','Webcam events','Reviewed','Visually confirmed','Accuracy'],1):
        summary.cell(4, col, name)
    order = list(dict.fromkeys([r[1] for r in rows]))
    selected = args.vp or VPS
    for vp in selected:
        if vp not in order:
            order.append(vp)
    end = max(6, 5+len(rows))
    for i, vp in enumerate(order, 5):
        summary.cell(i, 1, vp)
        summary.cell(i, 2, f'=COUNTIF(\'Blink Review\'!$B$6:$B${end},A{i})')
        summary.cell(i, 3, f'=COUNTIFS(\'Blink Review\'!$B$6:$B${end},A{i},\'Blink Review\'!$G$6:$G${end},0)+COUNTIFS(\'Blink Review\'!$B$6:$B${end},A{i},\'Blink Review\'!$G$6:$G${end},1)')
        summary.cell(i, 4, f'=COUNTIFS(\'Blink Review\'!$B$6:$B${end},A{i},\'Blink Review\'!$G$6:$G${end},1)')
        summary.cell(i, 5, f'=IF(OR(B{i}=0,C{i}<B{i}),"",D{i}/B{i})')
        summary.cell(i, 5).number_format = '0.0%'
    total = 5+len(order)
    summary.cell(total, 1, 'OVERALL')
    for col in (2,3,4):
        letter = 'BCDE'[col-2]
        summary.cell(total, col, f'=SUM({letter}5:{letter}{total-1})')
    summary.cell(total, 5, f'=IF(OR(B{total}=0,C{total}<B{total}),"",D{total}/B{total})')
    summary.cell(total, 5).number_format = '0.0%'
    for col, width in {'A':18,'B':21,'C':18,'D':25,'E':19}.items():
        summary.column_dimensions[col].width = width
    summary.freeze_panes='B5'
    summary.row_dimensions[1].height=31
    summary.row_dimensions[2].height=23
    summary.row_dimensions[4].height=29
    for rowno, bg in [(1,'102A43'),(2,'E8F1F8'),(4,'254D70'),(total,'DCEAF2')]:
        for cell in summary[rowno][:5]:
            cell.fill=PatternFill('solid', fgColor=bg)
            cell.font=Font(bold=rowno in (1,4,total),color='FFFFFF' if rowno in (1,4) else '263445',size=15 if rowno==1 else 10)
            cell.alignment=Alignment(vertical='center')
    if problems:
        err=wb.create_sheet('Warnings')
        err['A1']='Processing warnings — check before relying on complete counts'
        for i, message in enumerate(problems, 2): err.cell(i,1,message)
        err.column_dimensions['A'].width=100
    return wb


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--results-root',type=Path,default=Path('/Volumes/Empra10/Ergebnisse60'))
    parser.add_argument('--video-root',type=Path,default=Path('/Volumes/Empra9/Videos_60'))
    parser.add_argument('--vp',action='append')
    parser.add_argument('--trial',type=int,action='append')
    parser.add_argument('--reason-group',action='append',choices=['blink','suspected','invalid'],default=None)
    parser.add_argument('--output',type=Path,default=None)
    args=parser.parse_args()
    if args.reason_group is None:args.reason_group=['blink']
    rows,pictures,warnings=gather(args)
    if not rows:raise RuntimeError('No matching webcam blink events found; check --reason-group and input CSVs')
    target=args.output or args.results_root/'blink_extra_outputs'/'visual_review'/'webcam_blink_visual_review.xlsx'
    target.parent.mkdir(parents=True,exist_ok=True)
    print(f'[BUILD] {len(rows)} events; embedded images: {sum(bool(p) for p in pictures)}',flush=True)
    wb=make_book(rows,pictures,warnings,args)
    wb.save(str(target))
    print(f'[SAVED] {target}',flush=True)
    if warnings: print(f'[WARN] {len(warnings)} issues; see Warnings sheet',flush=True)

if __name__=='__main__':main()
